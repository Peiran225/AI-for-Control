"""Build high-resolution DOCX and PDF reports for the Transformer u(t) experiment."""

from __future__ import annotations

import os
import hashlib
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
FIG = ROOT / "paper_runs"
EQUATIONS = REPORTS / "equation_assets"
PY = os.environ.get("PYTHON", sys.executable)
RENDER = os.environ.get("RENDER_DOCX")


BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 41, 55)
MUTED = RGBColor(82, 96, 109)
HEADER_FILL = "F2F4F7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def clear_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_run_font(run, size: float | None = None, bold: bool = False, color: RGBColor | None = None, name: str = "Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc: Document, text: str = "", style: str | None = None, size: float = 10.5, bold: bool = False) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=DARK)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_run_font(r, size=8.3, color=MUTED)


def add_page_number_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    run = p.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 12.5, bold=True, color=BLUE)


def render_latex_equation(latex: str, number: int | str | None = None) -> Path | None:
    if not shutil.which("tectonic") or not shutil.which("pdftocairo"):
        return None
    EQUATIONS.mkdir(parents=True, exist_ok=True)
    key = hashlib.sha1(f"{number}:{latex}".encode("utf-8")).hexdigest()[:12]
    out_png = EQUATIONS / f"eq_{key}.png"
    if out_png.exists():
        return out_png

    if number is None:
        body = "\\[\n" f"{latex}\n" "\\]\n"
    else:
        body = (
            "\\begin{minipage}{6.15in}\n"
            "\\begin{equation}\n"
            f"\\tag{{{number}}}\n"
            f"{latex}\n"
            "\\end{equation}\n"
            "\\end{minipage}\n"
        )
    tex = (
        "\\documentclass[preview,border=2pt]{standalone}\n"
        "\\usepackage{amsmath,amssymb}\n"
        "\\begin{document}\n"
        f"{body}"
        "\\end{document}\n"
    )
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        tex_path = tmp_dir / "equation.tex"
        tex_path.write_text(tex, encoding="utf-8")
        subprocess.run(
            ["tectonic", "--outdir", str(tmp_dir), str(tex_path)],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        subprocess.run(
            ["pdftocairo", "-png", "-singlefile", "-r", "300", str(tmp_dir / "equation.pdf"), str(tmp_dir / "equation")],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        shutil.copyfile(tmp_dir / "equation.png", out_png)
    return out_png


def add_equation(doc: Document, latex: str, number: int | str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    png = render_latex_equation(latex, number)
    if png is not None:
        with Image.open(png) as im:
            width = min(6.35, max(2.2, im.width / 150.0))
        p.add_run().add_picture(str(png), width=Inches(width))
        return

    r = p.add_run(latex)
    set_run_font(r, size=10.5, name="Cambria Math")
    r.font.color.rgb = RGBColor(17, 24, 39)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 8.5) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        set_cell_width(hdr[i], widths[i])
        set_cell_shading(hdr[i], HEADER_FILL)
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color=DARK)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_width(cells[i], widths[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i >= 2 and i <= 5 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text)
            set_run_font(r, size=font_size, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_figure(doc: Document, rel_path: str, caption: str | None = None, width: float = 6.3) -> None:
    path = ROOT / rel_path
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        r = cp.add_run(caption)
        set_run_font(r, size=9, color=MUTED)


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    styles["Normal"].font.size = Pt(10.5)
    add_page_number_footer(doc)
    return doc


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=23, bold=True, color=DARK)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(12)
    sr = sp.add_run(subtitle)
    set_run_font(sr, size=10.5, color=MUTED)


def build_en() -> Path:
    doc = setup_doc()
    add_title(doc, "Transformer u(t) First-Layer Experiment Report", "Fixed-initial-condition time-dependent control trained by PMP/KKT optimality gaps")
    add_para(doc, "This report focuses on the manuscript's time-dependent control strategy u_theta(t): [0,T] -> [0,u_max] for the fixed initial condition. The state-dependent feedback extension u(t,N) is not included here.")

    add_heading(doc, "1. Model and Training Objective")
    add_para(doc, "We use the population dynamics and cost from the manuscript:")
    add_equation(doc, r"\dot N_i(t)=\big(r_i-\phi_i u(t)-M_iG(N(t))\big)N_i(t)", 1)
    add_equation(doc, r"G(N)=\log\left(1+\frac{1}{m}\sum_{k=1}^m N_k\right)", 2)
    add_equation(doc, r"J(u)=\alpha^\top N(T)+\int_0^T\big(\beta^\top N(t)+\gamma u(t)\big)\,dt", 3)
    add_para(doc, "Reported parameters: T=10, m=21, u_max=3, alpha=1, beta=0.1, gamma=20, and N_i(0)=10. The control is represented by a small Transformer encoder over normalized time.")
    add_para(doc, "Given u_theta(t), we roll out N_theta(t), solve the costate equation backward, and compute the switching function")
    add_equation(doc, r"\psi(t)=H_u(N,\lambda,u)=\gamma-\sum_i\phi_i\lambda_i(t)N_i(t)", 4)
    add_table(
        doc,
        ["component", "role"],
        [
            ["non-singular KKT loss", "enforces u=0 when psi>0 and u=u_max when psi<0"],
            ["singular loss", "near psi=0, matches u_theta(t) to the singular candidate u_sing(N(t))"],
        ],
        [1.7, 4.6],
        font_size=9.2,
    )
    add_para(doc, "Here q(t) is a smooth weight that switches between the singular condition near psi(t)=0 and the boundary KKT condition away from psi(t)=0.")

    add_heading(doc, "2. Learned u(t), N(t), and N(u)")
    add_para(doc, "The learned control starts high, transitions to a lower interior/singular-like region, and increases again near the terminal portion.")
    add_figure(
        doc,
        "paper_runs/open_loop_ut_report/ut_nt_trajectory_clean.png",
        caption="Figure 1. Learned Transformer control u(t), population trajectory N(t), switching function psi(t), and singular weight q(t).",
        width=6.35,
    )
    add_para(doc, "The N(u) phase plot below uses the total population across all 21 subpopulations.")
    add_figure(
        doc,
        "paper_runs/open_loop_ut_report/nu_phase_plot_clean.png",
        caption="Figure 2. N(u) phase plot using the total population across the 21 subpopulations.",
        width=6.75,
    )

    doc.add_page_break()
    add_heading(doc, "3. Training Loss Trajectories for PMP/KKT Conditions")
    add_para(doc, "The table reports the smallest recorded training loss. In this experiment, the training loss is the manuscript's PMP/KKT optimality gap, composed of the singular condition and the non-singular Hamiltonian minimization condition.")
    add_table(
        doc,
        ["metric", "value"],
        [
            ["total training loss / PMP-KKT optimality gap", "0.02637"],
            ["singular-condition training loss", "0.00167"],
            ["non-singular Hamiltonian-minimization training loss", "0.02471"],
            ["objective J on the training grid", "384.76"],
            ["range and mean of u_theta(t)", "min 1.038, max 2.758, mean 1.372"],
            ["terminal mean state", "1.207"],
        ],
        [3.2, 2.0],
        font_size=9.2,
    )
    add_figure(
        doc,
        "paper_runs/open_loop_ut_report/training_loss_trajectory_clean.png",
        caption="Figure 3. Training trajectories of the total PMP/KKT optimality gap and its singular and non-singular components.",
        width=6.35,
    )
    add_para(doc, "The final pointwise diagnostic below uses a smooth weight q(t) to separate the two regimes: near psi(t)=0 it emphasizes the singular-condition error; away from psi(t)=0 it emphasizes the boundary KKT error.")
    add_figure(
        doc,
        "paper_runs/open_loop_ut_report/pmp_condition_components_clean.png",
        caption="Figure 4. Pointwise singular-condition error and boundary KKT error along the final learned trajectory.",
        width=6.35,
    )

    add_heading(doc, "4. First-Layer Benchmark and Related Work")
    add_para(doc, "For this first-layer experiment, all numerical comparisons keep the same fixed initial condition and compare time-dependent controls u(t). Several related-work papers in the manuscript target value-function or feedback formulations; those are important method references, but they are not the same numerical task as this u(t) reproduction.")
    add_table(
        doc,
        ["reference", "learned object / formulation", "relation to this report"],
        [
            ["HJB / BSDE methods [1,2,7]", "value function V(t,N) or HJB PDE solution", "state-domain feedback/value formulation; not a direct u(t) baseline"],
            ["Neural-PMP [3]", "control sequence via forward rollout, backward costate recursion, and Hamiltonian-gradient updates", "closest related-work baseline for the present u(t) experiment; implemented numerically below"],
            ["DeepONet / PINN policy iteration [5,6]", "policy evaluation and improvement for HJB-type equations", "methodological reference for feedback/value learning; separate from this fixed-trajectory u(t) layer"],
            ["classical chemotherapy OC [4]", "PMP and singular-control structure", "source of the singular-control condition used in the manuscript"],
        ],
        [1.55, 2.3, 2.6],
        font_size=7.8,
    )
    doc.add_page_break()
    add_para(doc, "Among these, [3] is the closest apples-to-apples numerical comparison. Here [3] refers to Gu, Xiong, and Chen, Pontryagin Optimal Control via Neural Networks (arXiv:2212.14566). Their Neural-PMP method first learns a differentiable dynamics model and then updates a control sequence using PMP gradients. Since the dynamics are already known in the manuscript setting, we compare against the oracle-dynamics controller stage: forward state integration, backward costate recursion, and Hamiltonian-gradient updates of a discrete control sequence.")
    add_table(
        doc,
        ["method", "training / update criterion", "objective J", "J - direct", "PMP/KKT gap", "note"],
        [
            ["direct grid reference", "direct minimization of discretized J", "386.438", "0.000", "1.016", "cost reference"],
            ["Transformer u(t), 6 runs", "paper PMP/KKT optimality-gap loss", "386.738 +/- 0.045", "0.300 +/- 0.045", "0.463 +/- 0.148", "main reproduction"],
            ["best Transformer run", "same as above", "386.695", "0.257", "0.207", "best reported run"],
            ["Neural-PMP [3]", "Hamiltonian-gradient update with known dynamics", "386.986", "0.548", "8.574", "related-work baseline"],
            ["constant u=1.5", "fixed control", "400.403", "13.965", "76.556", "scale check"],
            ["repository demo output", "provided s.csv", "422.670", "36.232", "433.349", "original demo artifact"],
        ],
        [1.3, 1.95, 0.8, 0.8, 0.8, 1.0],
        font_size=7.35,
    )
    add_note(doc, "* Objective J is computed after fixing u(t), reintegrating N(t) with the same fine-step fourth-order Runge-Kutta evaluator, and applying the manuscript objective definition. The direct row is a numerical cost reference, not a neural model.")
    add_para(doc, "The Transformer runs are stable across random seeds and remain close to the direct cost reference. Compared with the Neural-PMP controller-stage baseline [3], the Transformer has both lower objective J and a smaller PMP/KKT diagnostic gap in this fixed-initial-condition u(t) experiment.")
    add_figure(
        doc,
        "paper_runs/first_layer_ut_benchmark/transformer_seed_loss_trajectories.png",
        caption="Figure 5. Multi-seed convergence of the Transformer u(t) PMP/KKT training loss.",
        width=6.25,
    )
    add_figure(
        doc,
        "paper_runs/first_layer_ut_benchmark/first_layer_objective_gap_closeup.png",
        caption="Figure 6. Objective gap relative to the direct cost reference for the main u(t) comparison.",
        width=6.25,
    )
    add_figure(
        doc,
        "paper_runs/neural_pmp_baseline_beta01/neural_pmp_ut_nt.png",
        caption="Figure 7. Neural-PMP controller-stage baseline [3]: control sequence and resulting state trajectory.",
        width=6.35,
    )
    add_figure(
        doc,
        "paper_runs/neural_pmp_baseline_beta01/neural_pmp_training_curve.png",
        caption="Figure 8. Neural-PMP controller-stage baseline [3]: selected-run training trajectories.",
        width=6.35,
    )

    add_heading(doc, "5. Conclusion")
    add_para(doc, "The requested first-layer u(t) reproduction is complete. The Transformer control trained with the manuscript's PMP/KKT optimality-gap loss is smooth, satisfies the control bounds, and reduces the training optimality gap from about 76 to 0.026 in the best run. Across six Transformer runs, the common-evaluator objective is 386.738 +/- 0.045, close to the direct cost reference 386.438 and better than the Neural-PMP [3] controller-stage baseline in this setting. The state-dependent extension u(t,N) is left for separate work because it requires different optimality conditions.")
    out = REPORTS / "ut_reproduction_report.docx"
    doc.save(out)
    return out


def build_zh() -> Path:
    doc = setup_doc()
    add_title(doc, "Transformer u(t) 第一层实验报告", "固定初始条件下的时间相关控制，使用 PMP/KKT optimality gap 训练")
    add_para(doc, "本报告只关注固定初始条件下的时间相关控制策略 u_theta(t): [0,T] -> [0,u_max]。不包含状态相关反馈控制 u(t,N)，因为该情形的最优性条件不同。")

    add_heading(doc, "1. 模型和训练目标")
    add_para(doc, "我们使用论文中的种群动力学模型和目标函数：")
    add_equation(doc, r"\dot N_i(t)=\big(r_i-\phi_i u(t)-M_iG(N(t))\big)N_i(t)", 1)
    add_equation(doc, r"G(N)=\log\left(1+\frac{1}{m}\sum_{k=1}^m N_k\right)", 2)
    add_equation(doc, r"J(u)=\alpha^\top N(T)+\int_0^T\big(\beta^\top N(t)+\gamma u(t)\big)\,dt", 3)
    add_para(doc, "本次参数为 T=10，m=21，u_max=3，alpha=1，beta=0.1，gamma=20，初始条件为 N_i(0)=10。控制函数由小型 Transformer encoder 在归一化时间上表示。")
    add_para(doc, "给定 u_theta(t) 后，先正向求解得到 N_theta(t)，再反向求解 costate，并计算 switching function：")
    add_equation(doc, r"\psi(t)=H_u(N,\lambda,u)=\gamma-\sum_i\phi_i\lambda_i(t)N_i(t)", 4)
    add_table(
        doc,
        ["component", "作用"],
        [
            ["non-singular KKT loss", "当 psi>0 时推动 u=0；当 psi<0 时推动 u=u_max"],
            ["singular loss", "当 psi 接近 0 时，使 u_theta(t) 接近 singular candidate u_sing(N(t))"],
        ],
        [1.85, 4.45],
        font_size=9.0,
    )
    add_para(doc, "其中 q(t) 是一个平滑权重，用来在 psi(t) 接近 0 的 singular condition 和 psi(t) 远离 0 的边界 KKT condition 之间切换。")

    add_heading(doc, "2. 学到的 u(t)、N(t) 和 N(u)")
    add_para(doc, "学到的开环控制在开始阶段较高，中间阶段下降到接近 singular control 的内部取值，末端附近再次升高。")
    add_figure(
        doc,
        "paper_runs/open_loop_ut_report/ut_nt_trajectory_clean.png",
        caption="图 1. 学到的 Transformer 控制 u(t)、状态轨迹 N(t)、switching function psi(t) 和 singular weight q(t)。",
        width=6.35,
    )
    add_para(doc, "下面的 N(u) 图以 21 个 subpopulation 的总和作为纵轴。")
    add_figure(
        doc,
        "paper_runs/open_loop_ut_report/nu_phase_plot_clean.png",
        caption="图 2. N(u) 相图，纵轴为 21 个 subpopulation 的总和。",
        width=6.75,
    )

    doc.add_page_break()
    add_heading(doc, "3. 各 PMP/KKT 最优性条件的 training loss trajectory")
    add_para(doc, "下表汇总训练过程中记录到的最小 training loss。这里的 training loss 正是论文中定义的 PMP/KKT optimality gap，由 singular condition 和 non-singular Hamiltonian minimization condition 两部分组成。")
    add_table(
        doc,
        ["metric", "value"],
        [
            ["total training loss / PMP-KKT optimality gap", "0.02637"],
            ["singular-condition training loss", "0.00167"],
            ["non-singular Hamiltonian-minimization training loss", "0.02471"],
            ["objective J on the training grid", "384.76"],
            ["range and mean of u_theta(t)", "min 1.038, max 2.758, mean 1.372"],
            ["terminal mean state", "1.207"],
        ],
        [3.2, 2.0],
        font_size=9.2,
    )
    add_para(doc, "训练曲线将论文中的两类最优性条件分开展示：当 psi(t) 接近 0 时，对应 singular condition；当 psi(t) 不为 0 时，对应 Hamiltonian 在控制边界 0 或 u_max 上取最小的 non-singular condition。结果显示，singular-condition gap 较小，主要剩余误差来自 non-singular Hamiltonian-minimization condition，尤其是控制切换附近。")
    add_figure(
        doc,
        "paper_runs/open_loop_ut_report/training_loss_trajectory_clean.png",
        caption="图 3. 总 PMP/KKT optimality gap 以及 singular、non-singular 两个组成部分的训练轨迹。",
        width=6.35,
    )
    add_para(doc, "下面的逐点诊断图使用平滑权重 q(t) 区分两类区域：当 psi(t) 接近 0 时，主要看 singular-condition error；当 psi(t) 远离 0 时，主要看边界 KKT error。")
    add_figure(
        doc,
        "paper_runs/open_loop_ut_report/pmp_condition_components_clean.png",
        caption="图 4. 最终学到的轨迹上的逐点 singular-condition error 和边界 KKT error。",
        width=6.35,
    )

    add_heading(doc, "4. 第一层 Benchmark 和 Related Work 对比")
    add_para(doc, "这一层实验固定初始条件，只比较时间相关控制 u(t)。论文 related work 里的若干方法面向 value function 或 feedback formulation，它们是重要的方法参考，但并不是和本报告完全相同的 u(t) 数值任务。")
    add_table(
        doc,
        ["引用", "学习对象 / formulation", "与本报告的关系"],
        [
            ["HJB / BSDE methods [1,2,7]", "value function V(t,N) 或 HJB PDE solution", "面向 state-domain feedback/value formulation；不是直接的 u(t) baseline"],
            ["Neural-PMP [3]", "forward rollout、backward costate recursion 和 Hamiltonian-gradient 更新控制序列", "与本报告 u(t) 实验最接近；下面给数值对比"],
            ["DeepONet / PINN policy iteration [5,6]", "HJB 型方程的 policy evaluation / improvement", "适合作为 feedback/value learning 的方法参考；属于另一层问题"],
            ["classical chemotherapy OC [4]", "PMP 和 singular-control 结构", "提供论文使用的 singular-control 条件"],
        ],
        [1.55, 2.3, 2.6],
        font_size=7.6,
    )
    doc.add_page_break()
    add_para(doc, "其中 [3] 是最接近同任务的数值比较。这里 [3] 指 Gu, Xiong, and Chen 的 Pontryagin Optimal Control via Neural Networks (arXiv:2212.14566)。该文 Neural-PMP 方法先学习可微动力学模型，再用 PMP gradient 更新控制序列。由于本报告中的动力学已知，我们比较的是其 oracle-dynamics controller stage：正向求解状态、反向求解 costate，并用 Hamiltonian gradient 更新离散控制序列。")
    add_table(
        doc,
        ["方法", "训练/更新准则", "目标函数 J", "J - direct", "PMP/KKT gap", "说明"],
        [
            ["direct grid reference", "直接最小化离散 J", "386.438", "0.000", "1.016", "成本参考"],
            ["Transformer u(t), 6 runs", "论文 PMP/KKT optimality-gap loss", "386.738 +/- 0.045", "0.300 +/- 0.045", "0.463 +/- 0.148", "主复现实验"],
            ["best Transformer run", "同上", "386.695", "0.257", "0.207", "最优 run"],
            ["Neural-PMP [3]", "已知动力学下的 Hamiltonian-gradient 控制更新", "386.986", "0.548", "8.574", "related-work baseline"],
            ["constant u=1.5", "固定控制", "400.403", "13.965", "76.556", "尺度参考"],
            ["repository demo output", "提供的 s.csv", "422.670", "36.232", "433.349", "原始 demo artifact"],
        ],
        [1.3, 1.95, 0.8, 0.8, 0.8, 1.0],
        font_size=7.25,
    )
    add_note(doc, "* 表中的 J 是在固定 u(t) 后，用同一个细步长四阶 Runge-Kutta evaluator 重新求解 N(t)，再按论文目标函数定义计算得到。direct 行是成本参考解，不是神经网络模型。")
    add_para(doc, "Transformer 在多个随机种子下结果稳定，并且接近 direct cost reference。和 Neural-PMP controller-stage baseline [3] 相比，在这个固定初始条件的 u(t) 实验中，Transformer 同时具有更低的目标函数 J 和更小的 PMP/KKT diagnostic gap。")
    add_figure(
        doc,
        "paper_runs/first_layer_ut_benchmark/transformer_seed_loss_trajectories.png",
        caption="图 5. Transformer u(t) 在多个随机种子下的 PMP/KKT training loss 收敛曲线。",
        width=6.25,
    )
    add_figure(
        doc,
        "paper_runs/first_layer_ut_benchmark/first_layer_objective_gap_closeup.png",
        caption="图 6. 主 u(t) 对比中，相对于 direct cost reference 的目标函数差值。",
        width=6.25,
    )
    add_figure(
        doc,
        "paper_runs/neural_pmp_baseline_beta01/neural_pmp_ut_nt.png",
        caption="图 7. Neural-PMP controller-stage baseline [3] 的控制序列和对应状态轨迹。",
        width=6.35,
    )
    add_figure(
        doc,
        "paper_runs/neural_pmp_baseline_beta01/neural_pmp_training_curve.png",
        caption="图 8. Neural-PMP controller-stage baseline [3] 的 selected-run 训练轨迹。",
        width=6.35,
    )

    add_heading(doc, "5. 结论")
    add_para(doc, "第一层 u(t) 复现实验已经完成。使用论文中的 PMP/KKT optimality-gap loss 训练得到的 Transformer 控制轨迹平滑并满足控制约束；best run 的 training optimality gap 从约 76 降到 0.026。六个 Transformer run 在同一 evaluator 下的目标函数为 386.738 +/- 0.045，接近 direct cost reference 386.438，并优于本设置下的 Neural-PMP [3] controller-stage baseline。本报告不包含状态相关扩展 u(t,N)。")
    out = REPORTS / "open_loop_ut_reproduction_report_zh.docx"
    doc.save(out)
    return out


def convert_to_pdf(docx_path: Path) -> Path:
    env = os.environ.copy()
    env["TMPDIR"] = "/private/tmp"
    env["TEMP"] = "/private/tmp"
    env["TMP"] = "/private/tmp"
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(REPORTS), str(docx_path)],
        check=True,
        cwd=str(ROOT),
        env=env,
    )
    return docx_path.with_suffix(".pdf")


def render(docx_path: Path) -> None:
    if not RENDER:
        print(f"Skipping page render for {docx_path.name}; set RENDER_DOCX to enable it.")
        return
    out_dir = REPORTS / f"{docx_path.stem}_render"
    subprocess.run(
        [PY, RENDER, str(docx_path), "--output_dir", str(out_dir), "--emit_pdf"],
        check=True,
        cwd=str(ROOT),
    )


def main() -> None:
    for builder in (build_en, build_zh):
        docx_path = builder()
        pdf_path = convert_to_pdf(docx_path)
        render(docx_path)
        print(docx_path)
        print(pdf_path)


if __name__ == "__main__":
    main()
